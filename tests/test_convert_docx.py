"""Markdown -> Word (.docx) through the multipart endpoint."""

from __future__ import annotations

import re

import httpx
import pytest
from docx.enum.section import WD_ORIENT
from fastapi.testclient import TestClient
from lxml import etree

from tests.utils import (
    DOCX_MEDIA_TYPE,
    W_NS,
    WORD_IMAGE_SIGNATURES,
    convert,
    docx_document,
    docx_drawing_count,
    docx_media,
    docx_paragraph_styles,
    docx_parts,
    docx_text,
    docx_xml,
    expect_ok,
    warnings_of,
)

MERMAID_DIAGRAMS_IN_COMPLETE_SAMPLE = 3


@pytest.fixture(scope="module")
def complete_docx(client: TestClient, complete_md: bytes) -> httpx.Response:
    return expect_ok(convert(client, "docx", complete_md, "complete.md"), DOCX_MEDIA_TYPE)


def _footer_parts(content: bytes) -> dict[str, str]:
    return {
        name: data.decode("utf-8")
        for name, data in docx_parts(content).items()
        if re.fullmatch(r"word/footer\d*\.xml", name)
    }


def test_opens_with_python_docx(complete_docx: httpx.Response) -> None:
    document = docx_document(complete_docx.content)
    assert document.paragraphs
    styles = docx_paragraph_styles(complete_docx.content)
    assert "Relatório de Testes" in styles.get("Title", [])
    assert "Visão geral" in styles.get("Heading 1", [])
    assert "Lista aninhada" in styles.get("Heading 2", [])


def test_complete_sample_has_no_warnings(complete_docx: httpx.Response) -> None:
    assert warnings_of(complete_docx) == []


def test_text_and_accents(complete_docx: httpx.Response) -> None:
    text = docx_text(complete_docx.content)
    for expected in ("coração", "informação", "Ünïcódé", "Nível três", "codigo_inline", "🚀"):
        assert expected in text, expected
    assert "segredo-do-comentario" not in text


def test_tables(complete_docx: httpx.Response) -> None:
    document = docx_document(complete_docx.content)
    assert len(document.tables) >= 2
    cells = [cell for table in document.tables for row in table.rows for cell in row.cells]
    texts = [cell.text for cell in cells]
    assert "Conversão" in texts
    assert "Gama" in texts
    assert any(re.search(r"linha 1\s+linha 2", text) for text in texts), "<br> lost in cell"
    code_runs = [
        run
        for cell in cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.text.strip() == "x = 1"
    ]
    assert code_runs and code_runs[0].style.name == "Verbatim Char"


def test_images_cover_every_mermaid_diagram(complete_docx: httpx.Response) -> None:
    # Three Mermaid diagrams plus one data: URI image.
    assert docx_drawing_count(complete_docx.content) >= MERMAID_DIAGRAMS_IN_COMPLETE_SAMPLE + 1
    media = docx_media(complete_docx.content)
    assert len(media) >= MERMAID_DIAGRAMS_IN_COMPLETE_SAMPLE + 1
    for name, data in media.items():
        assert data.startswith(WORD_IMAGE_SIGNATURES), f"{name} is not PNG/JPEG/GIF"
    assert "flowchart LR" not in docx_text(complete_docx.content)


def test_code_blocks_use_source_code_style(complete_docx: httpx.Response) -> None:
    code = "\n".join(docx_paragraph_styles(complete_docx.content).get("Source Code", []))
    assert "def calcular_total(itens):" in code
    assert "    return sum(itens)" in code  # indentation preserved
    assert "function somarValores(a, b) {" in code


def test_nested_list_levels(complete_docx: httpx.Response) -> None:
    xml = docx_xml(complete_docx.content)
    assert re.search(r'<w:ilvl w:val="3"\s*/>', xml), "4th list level (indentation) lost"


def test_math_is_native_word_equations(complete_docx: httpx.Response) -> None:
    xml = docx_xml(complete_docx.content)
    assert re.search(r"<m:oMath[\s>]", xml)  # inline $...$
    assert len(re.findall(r"<m:oMathPara[\s>]", xml)) >= 2  # $$...$$ and ```math


def test_footnotes(complete_docx: httpx.Response) -> None:
    assert "Texto da nota de rodapé." in docx_text(complete_docx.content, "word/footnotes.xml")


def test_task_list_checkboxes(complete_docx: httpx.Response) -> None:
    # pandoc draws the checkboxes as list bullets (numbering definitions).
    numbering = docx_xml(complete_docx.content, "word/numbering.xml")
    assert re.search(r'<w:lvlText w:val="☒"\s*/>', numbering)
    assert re.search(r'<w:lvlText w:val="☐"\s*/>', numbering)
    text = docx_text(complete_docx.content)
    assert "Tarefa concluída" in text
    assert "[x]" not in text and "[ ]" not in text


def _quote_boxes(content: bytes) -> list[tuple[str, str]]:
    """(bar color, text) of every quote box: a one-cell table with only a left border."""
    root = etree.fromstring(docx_parts(content)["word/document.xml"])
    boxes = []
    for table in root.iter(f"{{{W_NS}}}tbl"):
        left = table.find(f"{{{W_NS}}}tblPr/{{{W_NS}}}tblBorders/{{{W_NS}}}left")
        if left is None or left.get(f"{{{W_NS}}}val") != "single":
            continue
        text = " ".join("".join(table.itertext()).split())
        boxes.append((left.get(f"{{{W_NS}}}color", ""), text))
    return boxes


def test_alert_becomes_a_colored_box(complete_docx: httpx.Response) -> None:
    boxes = _quote_boxes(complete_docx.content)
    note = [text for color, text in boxes if color == "0969DA"]
    assert note, boxes
    # Bold title, translated (the sample is ``lang: pt-BR``), then the alert's text.
    assert note[0].startswith("Nota") and "Texto do alerta de nota." in note[0], note


def test_quotes_keep_lists_code_and_nesting_inside_the_box(client: TestClient) -> None:
    markdown = (
        "> [!TIP]\n> Rode o comando:\n>\n> ```bash\n> make build\n> ```\n>\n"
        "> - item um\n> - item dois\n\n"
        "> Nível 1\n>\n> > Nível 2\n> >\n> > > Nível 3\n"
    )
    content = expect_ok(convert(client, "docx", markdown), DOCX_MEDIA_TYPE).content
    boxes = _quote_boxes(content)
    tip = next(text for color, text in boxes if color == "1A7F37")
    for part in ("Rode o comando:", "make build", "item um", "item dois"):
        assert part in tip, (part, tip)
    # Nested quotes are nested boxes: each level's box contains the next ones.
    levels = [text for color, text in boxes if "Nível 1" in text or "Nível 3" in text]
    assert len(levels) == 3 and all("Nível 3" in text for text in levels), levels
    xml = docx_xml(content)
    # Every table cell ends with a paragraph (Word reports an error otherwise).
    assert not re.search(r"</w:tbl>\s*</w:tc>", xml)
    # Quote text keeps the muted "Block Text" style, outside the box's own layout.
    assert "Nível 1" in " ".join(docx_paragraph_styles(content).get("Block Text", []))


def test_inline_raw_html(complete_docx: httpx.Response) -> None:
    xml = docx_xml(complete_docx.content)
    assert re.search(r'<w:vertAlign w:val="superscript"\s*/>', xml), "<sup> lost"
    assert re.search(r'<w:vertAlign w:val="subscript"\s*/>', xml), "<sub> lost"
    text = docx_text(complete_docx.content)
    assert "Ctrl" in text
    assert "Conteúdo oculto do bloco details." in text
    assert "<kbd>" not in text and "<sup>" not in text


def test_default_section_is_a4_portrait_with_page_numbers(complete_docx: httpx.Response) -> None:
    section = docx_document(complete_docx.content).sections[-1]
    assert section.orientation == WD_ORIENT.PORTRAIT
    assert section.page_width.mm == pytest.approx(210, abs=1)
    assert section.page_height.mm == pytest.approx(297, abs=1)
    assert section.left_margin.mm == pytest.approx(20, abs=0.5)
    assert not section.footer.is_linked_to_previous
    footers = _footer_parts(complete_docx.content)
    assert any(re.search(r"\bPAGE\b", xml) for xml in footers.values()), footers


def test_landscape_letter_with_wide_margins(client: TestClient) -> None:
    response = convert(
        client,
        "docx",
        "# Paisagem\n\nTexto.",
        page_size="Letter",
        orientation="landscape",
        margin="wide",
    )
    section = docx_document(expect_ok(response, DOCX_MEDIA_TYPE).content).sections[-1]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width.mm == pytest.approx(279.4, abs=1)
    assert section.page_height.mm == pytest.approx(215.9, abs=1)
    assert section.left_margin.mm == pytest.approx(30, abs=0.5)
    assert section.right_margin.mm == pytest.approx(30, abs=0.5)


def test_without_page_numbers_there_is_no_footer(client: TestClient) -> None:
    response = convert(client, "docx", "# Sem rodapé\n\nTexto.", page_numbers=False)
    content = expect_ok(response, DOCX_MEDIA_TYPE).content
    assert docx_document(content).sections[-1].footer.is_linked_to_previous
    assert "w:footerReference" not in docx_xml(content)


def test_table_of_contents_field(client: TestClient, complete_md: bytes) -> None:
    response = convert(client, "docx", complete_md, toc=True)
    xml = docx_xml(expect_ok(response, DOCX_MEDIA_TYPE).content)
    assert re.search(r"TOC \\o", xml), "missing Word TOC field"


def test_latin1_input_is_decoded(client: TestClient, latin1_md: bytes) -> None:
    response = convert(client, "docx", latin1_md, "latin1.md")
    content = expect_ok(response, DOCX_MEDIA_TYPE).content
    assert "Relatório em Latin-1" in docx_paragraph_styles(content).get("Heading 1", [])
    assert "Coração, ação, informação" in docx_text(content)


def test_wide_tables_span_the_text_width(client: TestClient) -> None:
    long_cell = "texto bastante longo para não caber " * 2
    markdown = (
        "| Coluna | Descrição | Observações |\n|---|---|---|\n"
        f"| 1 | {long_cell} | {long_cell} |\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n"
    )
    xml = docx_xml(expect_ok(convert(client, "docx", markdown), DOCX_MEDIA_TYPE).content)
    widths = re.findall(r'<w:tblW w:type="(\w+)" w:w="(\d+)"\s*/>', xml)
    # Wide table: full text width (pct 5000 = 100 %); narrow table: its content's width
    # (LibreOffice and Google Docs would otherwise draw equal 1.375 in columns).
    assert [kind for kind, _ in widths] == ["pct", "pct"], widths
    assert widths[0][1] == "5000" and int(widths[1][1]) < 1500, widths
    grid = [int(width) for width in re.findall(r'<w:gridCol w:w="(\d+)"', xml)[:3]]
    assert grid[0] < grid[1] and grid[0] < grid[2], grid  # columns sized by their content


def test_adjacent_code_blocks_and_alerts_stay_apart(client: TestClient) -> None:
    markdown = (
        "```python\nprint(1)\n```\n\n```bash\necho 2\n```\n\n"
        "> [!NOTE]\n> Primeiro.\n\n> [!TIP]\n> Segundo.\n"
    )
    content = expect_ok(convert(client, "docx", markdown), DOCX_MEDIA_TYPE).content
    styles = [paragraph.style.name for paragraph in docx_document(content).paragraphs]
    # A borderless paragraph separates blocks that Word would otherwise merge into one box.
    for first, second in zip(styles, styles[1:], strict=False):
        assert not (first == second == "Source Code"), styles
    # Alerts are one-cell tables: two in a row need a paragraph between them.
    xml = docx_xml(content)
    assert xml.count('w:color="0969DA"') == 1 and xml.count('w:color="1A7F37"') == 1
    assert not re.search(r"</w:tbl>\s*<w:tbl>", xml)


def test_tables_with_many_columns_use_the_small_font(client: TestClient) -> None:
    header = [
        "ID",
        "Name",
        "Department",
        "Location",
        "Manager",
        "Status",
        "Active",
        "Salary",
        "Start",
        "Notes",
        "Extra",
    ]
    row = [
        "1",
        "Ana",
        "Engenharia",
        "São Paulo",
        "Bruno",
        "Ativo",
        "Sim",
        "12,500.00",
        "2024-01-01",
        "observação",
        "x",
    ]
    markdown = (
        f"| {' | '.join(header)} |\n|{'---|' * len(header)}\n| {' | '.join(row)} |\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n"
    )
    content = expect_ok(convert(client, "docx", markdown), DOCX_MEDIA_TYPE).content
    styles = docx_paragraph_styles(content)
    assert {"ID", "Department", "12,500.00"} <= set(styles.get("Compact Small", [])), styles
    assert "B" in styles.get("Compact", [])  # small tables keep the normal size
    xml = docx_xml(content)
    assert xml.count('<w:tblCellMar><w:left w:w="57" w:type="dxa"/>') == 1
