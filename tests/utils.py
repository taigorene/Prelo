"""Helpers shared by the test-suite: request shortcuts, PDF/DOCX inspection and sample builders."""

from __future__ import annotations

import io
import json
import mimetypes
import re
import struct
import unicodedata
import zipfile
import zlib
from collections.abc import Iterable
from typing import Any
from urllib.parse import unquote

import httpx
from docx import Document
from docx.document import Document as DocxDocument
from lxml import etree
from pypdf import PdfReader

from app.config import Settings

PDF_MEDIA_TYPE = "application/pdf"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# First bytes of /etc/passwd inside the container: must never leak into an output document.
PASSWD_MARKER = b"root:x:0:0"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_IMAGE_SIGNATURES = (b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"GIF87a", b"GIF89a")

# Deterministic settings for the suite: never depend on PRELO_* variables of the caller.
BASE_SETTINGS: dict[str, Any] = {
    "max_upload_mb": 20,
    "max_concurrent_jobs": 4,
    "conversion_timeout": 120,
    "zip_max_files": 1000,
    "zip_max_uncompressed_mb": 200,
    "allow_remote_images": True,
    "allow_private_networks": False,
    "remote_image_max_mb": 10,
    "remote_timeout": 5,
    "max_remote_images": 50,
    "cors_origins": "*",
    "root_path": "",
    "log_level": "warning",
}


def make_settings(**overrides: Any) -> Settings:
    """Build ``Settings`` from the suite defaults plus ``overrides`` (paths still come from env)."""
    return Settings(**{**BASE_SETTINGS, **overrides})


# --------------------------------------------------------------------------- requests


def _form_value(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def convert(
    client: httpx.Client,
    fmt: str,
    data: bytes | str,
    filename: str = "documento.md",
    **options: Any,
) -> httpx.Response:
    """POST ``data`` as a multipart upload to ``/api/v1/convert/{fmt}`` with form ``options``."""
    if isinstance(data, str):
        data = data.encode("utf-8")
    content_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    form = {key: _form_value(value) for key, value in options.items()}
    return client.post(
        f"/api/v1/convert/{fmt}", files={"file": (filename, data, content_type)}, data=form
    )


def describe(response: httpx.Response) -> str:
    """Short, readable summary of a response for assertion messages (never dumps binaries)."""
    content_type = response.headers.get("content-type", "")
    if content_type.startswith(("application/json", "text/")):
        body = response.text[:800]
    else:
        body = f"<{len(response.content)} bytes of {content_type or 'unknown type'}>"
    return f"HTTP {response.status_code}: {body}"


def expect_ok(response: httpx.Response, media_type: str) -> httpx.Response:
    """Assert a successful conversion of the given media type and return the response."""
    assert response.status_code == 200, describe(response)
    assert response.headers["content-type"].split(";")[0].strip() == media_type
    return response


def expect_error(response: httpx.Response, status: int, code: str | None = None) -> dict[str, Any]:
    """Assert the JSON error contract ``{"detail": str, "code": str}`` and return the body."""
    assert response.status_code == status, describe(response)
    body = response.json()
    assert isinstance(body.get("detail"), str), body
    assert body["detail"].strip(), body
    assert isinstance(body.get("code"), str), body
    if code is not None:
        assert body["code"] == code, body
    return body


def warnings_of(response: httpx.Response) -> list[str]:
    """Decode the ``X-Conversion-Warnings`` header (a JSON array of strings)."""
    raw = response.headers.get("x-conversion-warnings")
    assert raw is not None, "missing X-Conversion-Warnings header"
    value = json.loads(raw)
    assert isinstance(value, list), raw
    assert all(isinstance(item, str) for item in value), raw
    return value


def download_filename(response: httpx.Response) -> str:
    """File name announced by ``Content-Disposition`` (RFC 5987 ``filename*`` first)."""
    header = response.headers["content-disposition"]
    match = re.search(r"filename\*\s*=\s*UTF-8''([^;\s]+)", header, flags=re.IGNORECASE)
    if match:
        return unicodedata.normalize("NFC", unquote(match.group(1)))
    match = re.search(r'filename\s*=\s*"([^"]*)"', header, flags=re.IGNORECASE)
    assert match, header
    return match.group(1)


def ascii_filename(response: httpx.Response) -> str | None:
    """The plain ``filename="..."`` fallback of ``Content-Disposition``, if present."""
    match = re.search(r'filename\s*=\s*"([^"]*)"', response.headers["content-disposition"])
    return match.group(1) if match else None


# --------------------------------------------------------------------------- PDF


def normalize(text: str) -> str:
    """NFKC-normalize (folds ligatures and composes accents) and collapse whitespace."""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip()


def pdf_reader(content: bytes) -> PdfReader:
    assert content.startswith(b"%PDF-"), content[:32]
    return PdfReader(io.BytesIO(content))


def pdf_page_texts(content: bytes) -> list[str]:
    return [normalize(page.extract_text() or "") for page in pdf_reader(content).pages]


def pdf_text(content: bytes) -> str:
    return " ".join(pdf_page_texts(content))


def pdf_image_count(content: bytes) -> int:
    """Number of raster images (XObjects, including nested forms) over all pages."""
    return sum(len(page.images) for page in pdf_reader(content).pages)


def pdf_outline_titles(content: bytes) -> list[str]:
    def walk(items: Iterable[Any]) -> Iterable[str]:
        for item in items:
            if isinstance(item, list):
                yield from walk(item)
            else:
                yield normalize(str(item.title))

    return list(walk(pdf_reader(content).outline))


def pdf_page_size(content: bytes, index: int = 0) -> tuple[float, float]:
    box = pdf_reader(content).pages[index].mediabox
    return float(box.width), float(box.height)


# --------------------------------------------------------------------------- DOCX


def docx_document(content: bytes) -> DocxDocument:
    assert content[:4] == b"PK\x03\x04", content[:32]
    return Document(io.BytesIO(content))


def docx_parts(content: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def docx_xml(content: bytes, part: str = "word/document.xml") -> str:
    return docx_parts(content)[part].decode("utf-8")


def docx_media(content: bytes) -> dict[str, bytes]:
    parts = docx_parts(content)
    return {name: data for name, data in parts.items() if name.startswith("word/media/")}


def docx_text(content: bytes, part: str = "word/document.xml") -> str:
    """Text of every paragraph of a part (tables included), one paragraph per line."""
    root = etree.fromstring(docx_parts(content)[part])
    lines = []
    for paragraph in root.iter(f"{{{W_NS}}}p"):
        chunks = []
        for node in paragraph.iter(f"{{{W_NS}}}t", f"{{{W_NS}}}br", f"{{{W_NS}}}tab"):
            local = etree.QName(node).localname
            if local == "t":
                chunks.append(node.text or "")
            else:
                chunks.append("\n" if local == "br" else "\t")
        lines.append("".join(chunks))
    return unicodedata.normalize("NFC", "\n".join(lines))


def docx_drawing_count(content: bytes) -> int:
    root = etree.fromstring(docx_parts(content)["word/document.xml"])
    return sum(1 for _ in root.iter(f"{{{W_NS}}}drawing"))


def docx_all_paragraphs(document: DocxDocument) -> list[Any]:
    """Body paragraphs plus the paragraphs inside (nested) tables."""
    paragraphs = list(document.paragraphs)
    pending = list(document.tables)
    while pending:
        table = pending.pop()
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
                pending.extend(cell.tables)
    return paragraphs


def docx_paragraph_styles(content: bytes) -> dict[str, list[str]]:
    """Map of paragraph style name -> texts of the paragraphs using it."""
    styles: dict[str, list[str]] = {}
    for paragraph in docx_all_paragraphs(docx_document(content)):
        styles.setdefault(paragraph.style.name, []).append(paragraph.text)
    return styles


def assert_no_passwd_leak(content: bytes) -> None:
    """No part (or the raw bytes) of an output document may contain /etc/passwd."""
    assert PASSWD_MARKER not in content
    if content[:4] == b"PK\x03\x04":
        for name, data in docx_parts(content).items():
            assert PASSWD_MARKER not in data, f"/etc/passwd leaked into {name}"
    else:
        assert PASSWD_MARKER.decode() not in pdf_text(content)


# --------------------------------------------------------------------------- builders


def make_png(width: int = 4, height: int = 4, rgb: tuple[int, int, int] = (37, 99, 235)) -> bytes:
    """A valid solid-colour RGB PNG built without third-party libraries."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    raw = (b"\x00" + bytes(rgb) * width) * height
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def make_png_bomb(width: int = 20_000, height: int = 20_000) -> bytes:
    """A tiny PNG declaring a huge bitmap (1 bit per pixel; the pixel data is cut short)."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", width, height, 1, 0, 0, 0, 0)
    rows = zlib.compress((b"\x00" + b"\x00" * ((width + 7) // 8)) * 64)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", rows)


def pdf_text_sizes(content: bytes, word: str) -> list[float]:
    """Rendered size (pt, after all scaling) of the text lines containing ``word``.

    Chromium writes a line as several kerned runs; runs are grouped by their baseline.
    """
    lines: dict[tuple[int, float], list[str]] = {}
    sizes: dict[tuple[int, float], float] = {}
    for number, page in enumerate(pdf_reader(content).pages):

        def visit(text: str, cm: Any, tm: Any, font: Any, size: float, page: int = number) -> None:
            baseline = (page, round(cm[5] + cm[3] * tm[5], 1))
            lines.setdefault(baseline, []).append(text)
            sizes[baseline] = round(abs(size * tm[3] * cm[3]), 2)

        page.extract_text(visitor_text=visit)
    return [sizes[key] for key, texts in lines.items() if word in "".join(texts)]


# Served by the local image server of the SSRF tests (see ``conftest.py``).
PIXEL_PNG = make_png(6, 6, (220, 38, 38))

SVG_SAMPLE = (
    b'<svg xmlns="http://www.w3.org/2000/svg" width="120" height="60" viewBox="0 0 120 60">'
    b'<rect width="120" height="60" rx="8" fill="#0891b2"/>'
    b'<circle cx="30" cy="30" r="18" fill="#fff"/></svg>'
)

ZipEntry = tuple[str | zipfile.ZipInfo, bytes | str]


def build_zip(entries: Iterable[ZipEntry]) -> bytes:
    """Build an in-memory .zip; a ``ZipInfo`` key allows crafting hostile entries."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            info = name if isinstance(name, zipfile.ZipInfo) else zipfile.ZipInfo(name)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)
    return buffer.getvalue()


def mark_encrypted(zip_bytes: bytes) -> bytes:
    """Set the "encrypted" flag on every entry (local and central headers) of a .zip."""
    data = bytearray(zip_bytes)
    for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        start = data.find(signature)
        while start != -1:
            data[start + flag_offset] |= 0x01
            start = data.find(signature, start + 4)
    return bytes(data)
